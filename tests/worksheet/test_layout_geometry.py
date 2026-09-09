"""Rendered geometry regressions: a structural pass must not excuse wide gaps."""
import unittest,tempfile,sys,json
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
import pdfplumber
SKILL=Path(__file__).resolve().parents[2]/'skills/worksheet-dual-doc'
sys.path.insert(0,str(SKILL/'scripts'))
from inspect_input import inspect
from generate import generate
from common import office
from layout_audit import audit_question_layout

class GeometryTests(unittest.TestCase):
 def test_single_two_digit_and_reject_legacy_gap(self):
  with tempfile.TemporaryDirectory() as td:
   root=Path(td);src=root/'input.docx';d=Document();d.add_paragraph('高效作业20 探究世界的本质');d.add_paragraph('一、选择题')
   for n in range(1,13):
    d.add_paragraph(f'{n}. '+('观察事物要从实际出发，尊重客观规律。'*8)+'(B)')
    d.add_paragraph('A. 错误\tB. 正确\tC. 错误\tD. 错误');d.add_paragraph('【解析】故选B。')
   d.save(src);model=inspect(src);cfg=json.loads((SKILL/'assets/defaults.json').read_text());internal=root/'题目版.docx'
   manifest=generate(src,model,cfg,'题目版',internal);doc=office([internal],root/'doc','doc:MS Word 97')[0];pdf=office([doc],root/'pdf','pdf')[0]
   with pdfplumber.open(pdf) as pages:result=audit_question_layout(pages,model,manifest)
   self.assertEqual(result['errors'],[]);self.assertEqual(len(result['questions']),12)
   self.assertTrue(all(r['continuations_checked']>=2 for r in result['questions']))
   self.assertTrue(all(abs(r['gap_pt'])<1 for r in result['questions']))
   # Recreate the historic wide tab gap while preserving all question characters.
   d=Document(internal);p=next(p for p in d.paragraphs if p.text.startswith('('))
   p.paragraph_format.left_indent=Pt(63);p.paragraph_format.first_line_indent=Pt(-63)
   p.paragraph_format.tab_stops.add_tab_stop(Pt(63));p.runs[0]._r.append(OxmlElement('w:tab'));d.save(internal)
   bad=office([internal],root/'bad-doc','doc:MS Word 97')[0];pdf=office([bad],root/'bad-pdf','pdf')[0]
   with pdfplumber.open(pdf) as pages:result=audit_question_layout(pages,model,manifest)
   self.assertTrue(any('题号后间隔异常' in e for e in result['errors']),result)

if __name__=='__main__':unittest.main()
