import unittest,tempfile,sys,json,subprocess,shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches
from PIL import Image

SKILL=Path(__file__).resolve().parents[2]/'skills/worksheet-dual-doc'
sys.path.insert(0,str(SKILL/'scripts'))
from inspect_input import inspect
from selection import select
from generate import generate
from common import office,write_json,text
from verify_output import verify
from fonts import resolve_fonts

class PipelineTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
 def tearDown(self):self.tmp.cleanup()
 def fixture(self,count=1,separate=False,missing=False,image_before=False):
  d=Document();d.add_paragraph('高效作业20 探究世界的本质');d.add_paragraph('一、选择题')
  for n in range(1,count+1):
   if image_before and n==2:
    im=self.root/'picture.png';Image.new('RGB',(120,80),'blue').save(im);d.add_picture(str(im),width=Inches(1))
   d.add_paragraph(f'{n}.下面漫画体现的道理是'+('' if separate or missing else '( B )'))
   if separate:d.add_paragraph('( B )')
   d.add_paragraph('A. 不正确\tB. 正确\tC. 不符合\tD. 不符合')
   d.add_paragraph('【解析】 B符合题意。故选B。')
   d.add_paragraph('这是一段继续说明。')
  p=self.root/'input.docx';d.save(p);return p
 def test_missing_answer_rejected(self):
  m=inspect(self.fixture(missing=True));self.assertTrue(any('缺少明确答案' in x for x in m['errors']))
 def test_multiline_answer_and_two_digit_number(self):
  m=inspect(self.fixture(count=12,separate=True));self.assertEqual(m['errors'],[]);self.assertEqual(m['question_count'],12)
  self.assertEqual(m['questions'][-1]['answer'],'B');self.assertTrue(any(b['role']=='answer' and '继续说明' in b['text'] for b in m['blocks']))
 def test_preceding_cartoon_is_not_removed(self):
  m=inspect(self.fixture(count=2,image_before=True));self.assertEqual(m['errors'],[])
  images=[b for b in m['blocks'] if b.get('images')];self.assertEqual(images[0]['qid'],'q002');self.assertEqual(images[0]['role'],'body')
 def test_duplicate_number_rejected(self):
  p=self.fixture(count=2);d=Document(p)
  for x in d.paragraphs:
   if x.text.startswith('2.'):x.text=x.text.replace('2.','1.',1)
  d.save(p);self.assertIn('题号重复，不能自动生成',inspect(p)['errors'])
 def test_default_override_persistence(self):
  clone=self.root/'skill';shutil.copytree(SKILL,clone)
  script=clone/'scripts/worksheet.py';defaults=clone/'assets/defaults.json';before=defaults.read_bytes()
  r=subprocess.run([sys.executable,str(script),'config','--set','compiler=测试老师'],check=True,capture_output=True,text=True)
  self.assertEqual(json.loads(r.stdout)['compiler'],'测试老师');self.assertEqual(before,defaults.read_bytes())
  subprocess.run([sys.executable,str(script),'config','--set','compiler=测试老师','--save-defaults'],check=True,capture_output=True)
  self.assertEqual(json.loads(defaults.read_text())['compiler'],'测试老师')
 def test_real_font_auto_discovery(self):self.assertTrue(Path(resolve_fonts()['font_file']).exists())
 def test_final_doc_tampering_is_detected(self):
  src=self.fixture();m=select(inspect(src));cfg=json.loads((SKILL/'assets/defaults.json').read_text());cfg['compiler']='临时编制'
  run=self.root/'run';(run/'internal').mkdir(parents=True);(run/'candidates').mkdir()
  manifests={v:generate(src,m,cfg,v,run/'internal'/f'{v}.docx') for v in ['题目版','答案版']}
  write_json(run/'input.json',m);write_json(run/'manifest.json',{'config':cfg,'variants':manifests,'review_issues':[]})
  office(list((run/'internal').glob('*.docx')),run/'candidates','doc:MS Word 97')
  self.assertEqual(verify(run)['errors'],[])
  # Mutate saved candidate through DOCX then re-export, without changing the oracle.
  back=office([run/'candidates/题目版.doc'],run/'mutation','docx')[0];d=Document(back)
  for p in d.paragraphs:
   if '下面漫画' in p.text:p.text=p.text.replace('下面漫画','内容已被篡改');break
  d.save(back);office([back],run/'candidates','doc:MS Word 97')
  self.assertTrue(any('文字不一致' in e for e in verify(run)['errors']))


class BoundaryTests(unittest.TestCase):
 setUp=PipelineTests.setUp
 tearDown=PipelineTests.tearDown
 fixture=PipelineTests.fixture
 def test_table_spans_survive_doc_roundtrip(self):
  from common import table_structure
  src=self.fixture();d=Document(src);table=d.add_table(rows=45,cols=3)
  for i,row in enumerate(table.rows):
   for j,c in enumerate(row.cells):c.text=f'材料{i+1}列{j+1}：保留表格内容与合并关系。'
  table.cell(0,0).merge(table.cell(0,1));table.cell(1,0).merge(table.cell(2,0))
  d.save(src);m=inspect(src);cfg=json.loads((SKILL/'assets/defaults.json').read_text())
  out=self.root/'table.docx';generate(src,m,cfg,'答案版',out)
  binary=office([out],self.root/'doc','doc:MS Word 97')[0]
  reopened=office([binary],self.root/'back','docx')[0]
  self.assertEqual(table_structure(Document(src).tables[0]._tbl),table_structure(Document(reopened).tables[0]._tbl))
 def test_paragraph_tab_definitions_not_text_atoms(self):
  from common import patch_text
  from docx.shared import Pt
  d=Document();p=d.add_paragraph('12.保留题干(B)');p.paragraph_format.tab_stops.add_tab_stop(Pt(50))
  patch_text(p._p,0,3,'');self.assertEqual(text(p._p),'保留题干(B)')
 def test_painting_attachment(self):
  p=self.fixture(count=2,image_before=True);d=Document(p)
  for par in d.paragraphs:
   if par.text.startswith('2.'):par.text=par.text.replace('下面漫画','下面绘画作品')
  d.save(p);m=inspect(p);self.assertEqual(m['errors'],[])
  self.assertEqual(next(b for b in m['blocks'] if b.get('images'))['qid'],'q002')

if __name__=='__main__':unittest.main(verbosity=2)
