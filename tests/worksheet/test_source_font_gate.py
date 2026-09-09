"""Source font designation is required; local fallback must be disclosed."""
import unittest,tempfile,sys,json,shutil
from pathlib import Path
from unittest.mock import patch
from docx import Document
SKILL=Path(__file__).resolve().parents[2]/'skills/worksheet-dual-doc'
sys.path.insert(0,str(SKILL/'scripts'))
from inspect_input import inspect
from selection import select
from generate import generate
from common import office,write_json
from deletion_report import build
from verify_output import verify
from fonts import resolve_fonts

class SourceFontGateTests(unittest.TestCase):
 def test_source_fallback_allowed_with_explicit_warning(self):
  with tempfile.TemporaryDirectory() as td:
   root=Path(td);src=root/'input.docx';d=Document();d.add_paragraph('高效作业20 探究世界的本质');d.add_paragraph('一、选择题(本大题共2小题,每小题2分,共4分)')
   for n,t in enumerate(['2025·浙江1月选考应删除的题目','2025·温州二模'+('保留题干和正确的选项。'*12)],1):
    d.add_paragraph(f'{n}.{t}(A)');d.add_paragraph('A. 正确\tB. 错误\tC. 错误\tD. 错误');d.add_paragraph('【解析】故选A。')
   d.add_paragraph('二、综合题(8分)');d.add_paragraph('3.2025·浙江6月选考应删除的材料');d.add_paragraph('【答案】应删除答案');d.save(src)
   model=select(inspect(src));cfg=json.loads((SKILL/'assets/defaults.json').read_text());run=root/'run';(run/'internal').mkdir(parents=True);manifests={}
   for v in ['题目版','答案版']:manifests[v]=generate(src,model,cfg,v,run/'internal'/f'{v}.docx')
   write_json(run/'input.json',model);write_json(run/'manifest.json',{'config':cfg,'variants':manifests,'review_issues':[]})
   build(model,run/'candidates/删除记录.pdf','授权预览替代字体')
   # Deliberately permit a directory without LiSu in this negative-only fixture.
   # Only source spans may use fallback; body font checks remain strict.
   font_info=resolve_fonts();isolated=root/'fonts';isolated.mkdir()
   shutil.copy2(font_info['font_file'],isolated/Path(font_info['font_file']).name)
   kai=Path(font_info['directory'])/'Kaiti.ttf'
   if kai.exists():shutil.copy2(kai,isolated/kai.name)
   font_info={**font_info,'directory':str(isolated)}
   with patch('fonts.resolve_source_font',return_value=font_info),patch('fonts.resolve_fonts',return_value=font_info):
    office([run/'internal'/f'{v}.docx' for v in manifests],run/'candidates','doc:MS Word 97')
    result=verify(run)
   self.assertEqual(result['status'],'awaiting_visual_review')
   self.assertTrue(any('本机预览使用替代字体' in e for e in result['warnings']))
   other=[e for e in result['errors'] if not any(t in e for t in ['渲染字体','真实隶书'])]
   self.assertEqual(other,[],result['errors'])
   self.assertEqual(result['files']['题目版']['pages'],2)
   self.assertEqual(result['files']['答案版']['pages'],2)

if __name__=='__main__':unittest.main()
