"""Verify point geometry in the actual saved DOC, including long-point exceptions."""
import unittest,tempfile,sys,json
from pathlib import Path
from docx import Document
import pdfplumber
SKILL=Path(__file__).resolve().parents[2]/'skills/worksheet-dual-doc'
sys.path.insert(0,str(SKILL/'scripts'))
from inspect_input import inspect
from generate import generate
from common import office,norm
from layout_audit import audit_point_layout

class PointLayoutTests(unittest.TestCase):
 def test_saved_doc_short_long_split_paragraphs_and_negative_wrap(self):
  with tempfile.TemporaryDirectory() as td:
   root=Path(td);src=root/'input.docx';d=Document();d.add_paragraph('高效作业28 发展中国特色社会主义文化');d.add_paragraph('一、选择题')
   samples=[['①'+ '甲'*18,'②'+'乙'*18,'③'+'丙'*18,'④'+'丁'*18],['①'+'甲'*32,'②'+'乙'*18,'③'+'丙'*18,'④'+'丁'*18],['①'+'甲'*70,'②'+'乙'*18,'③'+'丙'*18,'④'+'丁'*18],['①'+'甲'*25,'②'+'乙'*10,'③'+'丙'*25,'④'+'丁'*10]]
   for n,items in enumerate(samples,1):
    d.add_paragraph(f'{n}.以下说法正确的是(B)')
    if n==1:
     for item in items:d.add_paragraph(item)
    else:d.add_paragraph(' '.join(items))
    d.add_paragraph('A.①②\tB.①③\tC.②④\tD.③④');d.add_paragraph('【解析】故选B。')
   d.save(src);model=inspect(src);self.assertEqual(model['errors'],[])
   cfg=json.loads((SKILL/'assets/defaults.json').read_text());internal=root/'题目版.docx';manifest=generate(src,model,cfg,'题目版',internal)
   def render(folder):
    saved=office([internal],root/folder,'doc:MS Word 97')[0]
    return office([saved],root/(folder+'-pdf'),'pdf')[0]
   with pdfplumber.open(render('good')) as pdf:result=audit_point_layout(pdf,model)
   self.assertEqual(result['errors'],[],result)
   qs={q['number']:q['points'] for q in result['questions']}
   self.assertEqual([p['lines'] for p in qs[1]],[1,1,1,1])
   self.assertAlmostEqual(qs[1][1]['x'],qs[1][3]['x'],places=1)
   self.assertEqual(qs[2][0]['x'],qs[2][1]['x'])
   self.assertGreater(qs[3][0]['lines'],1)
   self.assertAlmostEqual(qs[4][1]['x'],qs[4][3]['x'],places=1)
   self.assertEqual([m['point_items'] for m in manifest if 'point_items' in m],samples)
   # Restore the old free-flow text: content is identical, but one point wraps needlessly.
   d=Document(internal);p=next(p for p in d.paragraphs if p.text.startswith('①'));p.text=' '.join(samples[0]);d.save(internal)
   with pdfplumber.open(render('bad')) as pdf:bad=audit_point_layout(pdf,model)
   self.assertTrue(any('可整行容纳却跨行' in e for e in bad['errors']),bad)

if __name__=='__main__':unittest.main()
