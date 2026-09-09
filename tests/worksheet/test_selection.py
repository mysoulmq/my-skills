import unittest,tempfile,sys,json
from pathlib import Path
from docx import Document
from PIL import Image
SKILL=Path(__file__).resolve().parents[2]/'skills/worksheet-dual-doc'
sys.path.insert(0,str(SKILL/'scripts'))
from inspect_input import inspect
from selection import select,source_label
from generate import generate

class SelectionTests(unittest.TestCase):
 def test_delete_and_placeholder_keep_source_number(self):
  with tempfile.TemporaryDirectory() as td:
   src=Path(td)/'a.docx';d=Document();d.add_paragraph('高效作业20 探究世界的本质');d.add_paragraph('一、选择题(本大题共3小题,每小题2分,共6分)')
   for n,t in enumerate(['2025·浙江1月选考待删除的题干','2025·浙江二模应保留的题干','2026年浙江开展活动，选考安排属于正文'],1):
    d.add_paragraph(f'{n}.{t}(A)')
    if n==1:
     pic=Path(td)/'removed.png';Image.new('RGB',(31,31),'magenta').save(pic);d.add_picture(str(pic))
    d.add_paragraph('A. 正确\tB. 错误\tC. 错误\tD. 错误');d.add_paragraph('【解析】故选A。')
   d.add_paragraph('二、综合题(8分)');d.add_paragraph('4.2025·浙江6月选考需要删除的材料');d.add_paragraph('【答案】删除答案独有文字');d.save(src)
   m=select(inspect(src));self.assertFalse(m['errors']);self.assertEqual([q['number'] for q in m['questions']],[1,2,4]);self.assertEqual([r['source_number'] for r in m['removed_questions']],[1,4]);self.assertEqual(m['questions'][-1]['placeholder'],True)
   cfg=json.loads((SKILL/'assets/defaults.json').read_text())
   for version in ['题目版','答案版']:
    out=Path(td)/(version+'.docx');manifest=generate(src,m,cfg,version,out);text=''.join(p.text for p in Document(out).paragraphs)
    self.assertEqual(len(Document(out).inline_shapes),0);self.assertNotIn('待删除的题干',text);self.assertNotIn('删除答案独有文字',text);self.assertIn('(\u3000 \u3000)1.（2025·浙江二模）应保留',text);self.assertIn('2026年浙江开展活动',text);self.assertIn('共2小题',text);self.assertIn('共4分',text);self.assertNotIn('8分',text)
    self.assertEqual([b['text'] for b in manifest if b['role']=='placeholder'],['4.'])
 def test_label_boundaries(self):
  self.assertEqual(source_label('2.2023·河北高考大海边,浪拍打海岸(A)')['text'],'2023·河北高考')
  self.assertEqual(source_label('1.（2025·温州二模）生成式人工智能(A)')['text'],'2025·温州二模')
  self.assertIsNone(source_label('1.2026年是工程通水12周年(A)'))
  with self.assertRaises(ValueError):source_label('1.2026·地点未明确考试名正文')

if __name__=='__main__':unittest.main()
