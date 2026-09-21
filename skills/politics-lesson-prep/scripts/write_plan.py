"""Teacher-first DOCX: linked navigation, period cards, then detailed teaching cards."""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

LABELS={'ask':'教师提问','expected':'预期回答','misconception':'典型错误','followup':'追问纠正','explanation':'关键讲解','check':'理解检查','transition':'过渡'}
NAVY='173E52';TEAL='007F86';RED='A63C32'


def shade(cell, color):
    el=OxmlElement('w:shd');el.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(el)


def columns(table,widths):
    for col,width in zip(table.columns,widths):col.width=Cm(width)


def keep_row(row):
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))


def bookmark(par,name,number):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number));par._p.insert(0,start);par._p.append(end)


def link(par,label,target):
    h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),target)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),TEAL);props.append(color)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');props.append(u);run.append(props)
    t=OxmlElement('w:t');t.text=label;run.append(t);h.append(run);par._p.append(h)


def box(doc,label,value,color='EAF4F4'):
    table=doc.add_table(rows=1,cols=1);cell=table.cell(0,0);shade(cell,color)
    keep_row(table.rows[0])
    p=cell.paragraphs[0];r=p.add_run(label+'  ');r.bold=True;r.font.color.rgb=RGBColor.from_string(RED if color=='FFF0E9' else NAVY)
    p.add_run(value);doc.add_paragraph().paragraph_format.space_after=Pt(0)


def pages_for(a,mapping):return [m for m in mapping if a['id'] in m.get('activityIds',[])]


def page_label(pages):
    groups=[]
    for page in sorted(set(pages)):
        if groups and page==groups[-1][-1]+1:groups[-1].append(page)
        else:groups.append([page])
    return '、'.join('P'+str(g[0])+(f'—{g[-1]}' if len(g)>1 else '') for g in groups)


def newpage_heading(doc,text):
    p=doc.add_heading(text,1);p.paragraph_format.page_break_before=True;return p


def write(plan, mapping, output):
    if Path(output).exists():raise ValueError('Refusing to overwrite output')
    doc=Document();section=doc.sections[0]
    section.top_margin=section.bottom_margin=Cm(1.6);section.left_margin=section.right_margin=Cm(1.8)
    for border in doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    for name in ('Normal','Title','Heading 1','Heading 2','Heading 3'):
        style=doc.styles[name];style.font.name='Microsoft YaHei'
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
        style.font.size=Pt({'Normal':10,'Title':24,'Heading 1':17,'Heading 2':13,'Heading 3':11}[name])
        style.font.color.rgb=RGBColor.from_string('222222' if name=='Normal' else NAVY)
        style.paragraph_format.space_after=Pt(5);style.paragraph_format.space_before=Pt(0 if name in ('Normal','Title') else 7)
    doc.styles['Normal'].paragraph_format.line_spacing=1.12
    footer=section.footer.paragraphs[0];footer.alignment=2
    footer.add_run('授课方案 · ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.add_heading(plan['lesson'],0);doc.add_heading('授课方案｜先看速览，再按需查讲法',1)
    box(doc,'教学主线',plan.get('designRationale','').split('，',1)[0])
    doc.add_heading('目录 · 点击跳转',1)
    bookmark(doc.paragraphs[-1],'contents',1)
    for i,period in enumerate(plan['periods'],1):
        p=doc.add_paragraph();link(p,period['title'],f'period_{i}');p.add_run('    ');link(p,'详细讲法',f'detail_{i}')
    doc.add_heading('本课要达成什么',2)
    for item in plan['goals']:doc.add_paragraph('• '+item)
    box(doc,'重点突破','；'.join(plan['difficulties']),'FFF0E9')
    doc.add_paragraph('课前：看主线和每课时速览。课堂：使用PPT备注中的讲授卡。遇到难点：从目录进入详细讲法。底色用于区分教学重点与易错提醒。')
    for i,period in enumerate(plan['periods'],1):
        heading=newpage_heading(doc,period['title']);bookmark(heading,f'period_{i}',10+i)
        p=doc.add_paragraph();link(p,'返回目录','contents');p.add_run('   ');link(p,'本课时详细讲法',f'detail_{i}')
        card=period.get('quickCard',{})
        box(doc,'这一课怎么推进',card.get('mainline',''))
        doc.add_heading('必讲透的区别',2)
        for value in card.get('mustExplain',[]):
            p=doc.add_paragraph();p.add_run('● ').bold=True;p.add_run(value)
        doc.add_heading('40分钟课堂路线',2)
        table=doc.add_table(rows=1,cols=3);table.autofit=False;columns(table,[2.4,5,10])
        for cell,width,value in zip(table.rows[0].cells,[2.4,5,10],['时间 / 页码','活动','关键提问']):
            cell.width=Cm(width);cell.text=value;shade(cell,NAVY)
            for run in cell.paragraphs[0].runs:run.font.color.rgb=RGBColor(255,255,255);run.bold=True
        elapsed=0
        for j,a in enumerate(period['activities']):
            related=pages_for(a,mapping)
            if not related:raise ValueError(f'No actual slide mapping for activity {a["id"]}')
            row=table.add_row().cells
            row[0].text=f'{elapsed}—{elapsed+a["minutes"]}分\n'+page_label(m['page'] for m in related);elapsed+=a['minutes']
            row[1].text=a['title'];row[2].text=a.get('cue',{}).get('ask',a['teaching']['ask'])
            for cell in row:
                if j%2==0:shade(cell,'F0F5F7')
                for p in cell.paragraphs:p.paragraph_format.space_after=Pt(4)
            keep_row(table.rows[-1])
        doc.add_paragraph();box(doc,'时间取舍',card.get('timeChoice','优先核心推导；延伸练习按课堂理解情况安排。'),'FFF0E9')
    newpage_heading(doc,'教师课前掌握 · 按需查阅')
    box(doc,'为什么这样安排',plan.get('designRationale',''))
    for value in plan.get('preparation',[]):box(doc,'判断要点',value)
    shown_clicks=set()
    for i,period in enumerate(plan['periods'],1):
        heading=newpage_heading(doc,period['title']+'｜详细讲法');bookmark(heading,f'detail_{i}',100+i)
        p=doc.add_paragraph();link(p,'返回本课时速览',f'period_{i}')
        for a in period['activities']:
            doc.add_heading(a['title'],2);related=pages_for(a,mapping)
            doc.add_paragraph('用时 '+str(a['minutes'])+'分钟  ·  课件 '+page_label(m['page'] for m in related))
            cue=a.get('cue',{})
            box(doc,'讲解抓手',cue.get('explain',a['teaching']['explanation']))
            table=doc.add_table(rows=0,cols=2);table.autofit=False;columns(table,[2.3,15.1])
            for key in ('ask','expected','followup','explanation','check','transition'):
                row=table.add_row().cells;row[0].width=Cm(2.3);row[1].width=Cm(15.1)
                row[0].text=LABELS[key];shade(row[0],'EAF4F4');row[0].paragraphs[0].runs[0].bold=True
                row[1].text=a['teaching'][key]
                for cell in row:
                    for p in cell.paragraphs:p.paragraph_format.space_after=Pt(4)
                keep_row(table.rows[-1])
            doc.add_paragraph();box(doc,'易错与纠正',a['teaching']['misconception']+'\n'+cue.get('followup',a['teaching']['followup']),'FFF0E9')
            for m in related:
                steps=m.get('clickTexts',[])
                if not steps:continue
                if m['page'] in shown_clicks:continue
                shown_clicks.add(m['page'])
                p=doc.add_paragraph();p.add_run(f'P{m["page"]} · 点击顺序  ').bold=True
                p.add_run(' → '.join(f'{n} '+(' / '.join(values)[:36]+('…' if len(' / '.join(values))>36 else '')) for n,values in enumerate(steps,1)))
    if plan.get('timeCuts'):
        doc.add_heading('跨课时机动取舍',1)
        for value in plan['timeCuts']:doc.add_paragraph(value)
    Path(output).parent.mkdir(parents=True,exist_ok=True);doc.save(output)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('plan');p.add_argument('mapping');p.add_argument('output')
    a=p.parse_args();write(json.loads(Path(a.plan).read_text()),json.loads(Path(a.mapping).read_text()),a.output)
