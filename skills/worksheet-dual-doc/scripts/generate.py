from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from common import ROOT,text,norm,clone_block,patch_text,prune,image_hashes
from inspect_input import STEM,ANSWER

def choice_indent(number):
    # Reference prefix '(fullwidth-space, halfwidth-space, fullwidth-space)N.'.
    # SimSun 10.5pt: ASCII = 5.25pt; fullwidth space = 10.5pt.
    return 42.0 + 5.25 * len(str(number))

def style_run(r, color='000000',size=10.5,font='宋体'):
    rp=r.find(qn('w:rPr'))
    if rp is None:rp=OxmlElement('w:rPr');r.insert(0,rp)
    for name in ['rFonts','sz','szCs','color','vanish','webHidden','rStyle','spacing','w','position']:
        for old in list(rp.findall(qn('w:'+name))):rp.remove(old)
    f=OxmlElement('w:rFonts')
    for attr in ['ascii','hAnsi','eastAsia','cs']:f.set(qn('w:'+attr),font)
    rp.append(f)
    for name,val in [('sz',str(int(size*2))),('szCs',str(int(size*2))),('color',color)]:
        v=OxmlElement('w:'+name);v.set(qn('w:val'),val);rp.append(v)

def format_block(e,doc,role,choice=False,number=1):
    left=choice_indent(number)
    for p in [e] if e.tag==qn('w:p') else e.xpath('.//w:p'):
        pp=p.find(qn('w:pPr'))
        if pp is not None:p.remove(pp)
        par=Paragraph(p,doc._body);par.style=doc.styles['WorksheetAnswer' if role=='answer' else 'WorksheetBody']
        fmt=par.paragraph_format;fmt.left_indent=Pt(0);fmt.first_line_indent=Pt(0);fmt.line_spacing=Pt(14);fmt.line_spacing_rule=WD_LINE_SPACING.AT_LEAST;fmt.space_after=Pt(2);fmt.widow_control=True
        fmt.keep_with_next=role in ['section','stem'] or bool(e.xpath('.//w:drawing'));fmt.keep_together=False
        if role=='section':fmt.space_before=Pt(5)
        if choice and role in ['stem','body'] and e.tag==qn('w:p'):
            fmt.left_indent=Pt(left);fmt.first_line_indent=Pt(-left if role=='stem' else 0)
            for x in [left+100,left+200,left+300]:fmt.tab_stops.add_tab_stop(Pt(x))
        for r in p.xpath('.//w:r'):style_run(r,'FF0000' if role=='answer' else '000000')
    for row in e.xpath('.//w:tr'):
        trpr=row.find(qn('w:trPr'))
        if trpr is None:trpr=OxmlElement('w:trPr');row.insert(0,trpr)
        if trpr.find(qn('w:cantSplit')) is None:trpr.append(OxmlElement('w:cantSplit'))
        for h in row.xpath('./w:trPr/w:trHeight'):h.getparent().remove(h)
    for ext in e.xpath('.//wp:extent'):
        cx,cy=int(ext.get('cx')),int(ext.get('cy'));maxw=int((doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin)-Pt(left if choice and e.tag==qn('w:p') else 0))
        if cx>maxw:
            ratio=maxw/cx;ext.set('cx',str(maxw));ext.set('cy',str(int(cy*ratio)))
            drawing=ext.getparent()
            for x in drawing.xpath('.//a:xfrm/a:ext'):x.set('cx',str(maxw));x.set('cy',str(int(cy*ratio)))


def title_for(raw,config):
    if config.get('title'):return config['title']
    t=re.sub(r'^高效作业\s*\d+\s*','',raw)
    t=re.sub(r'[（(]见学生用书.*$','',t).strip()
    # Explicit content-topic mapping for the supplied textbook, not assignment-number guessing.
    lessons={'探究世界的本质':2,'把握世界的规律':3,'探索认识的奥秘':4,'寻觅社会的真谛':5,'实现人生的价值':6,'继承发展中华优秀传统文化':7}
    for topic,n in lessons.items():
        if t.startswith(topic):return f'第{n}课  {t}'
    return t

def generate(source,model,config,variant,out):
    doc=Document(ROOT/'assets/template.docx');src=Document(source)
    for name in ['WorksheetBody','WorksheetAnswer']:
        if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        st=doc.styles[name];st.font.name='宋体';st.font.size=Pt(10.5)
    # Position the logo without text wrapping; explicit paragraph indents provide
    # its reserved area consistently in WPS and LibreOffice.
    for anchor in doc._element.xpath('//wp:anchor'):
        for node in list(anchor):
            if node.tag.rsplit('}',1)[-1].startswith('wrap'):
                replacement=OxmlElement('wp:wrapNone');anchor.replace(node,replacement)
    # Template body contains only logo paragraph. Replace its old paragraph placement.
    p=doc.paragraphs[0];p.alignment=1;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Pt(45)
    p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_with_next=True
    title=title_for(model['title'],config)
    if len(title)>25 and '——' in title:title=title.replace('——','\n——',1)
    r=p.add_run(title);r.bold=True;style_run(r._r,size=14,font='华文楷体')
    # Nature reference: visible field starts at 117.50, 196.25, 290.75,
    # 369.50 and 443.00 pt from page edge. Margin is 45.35 pt.
    # Absolute reference coordinates must not be reused as relative indents.
    extra_comp=max(0, len(config['compiler'])-3)*10.5
    extra_proof=max(0, len(config['proofreader'])-3)*10.5
    meta=doc.add_paragraph(f"编制人：{config['compiler']}\t校对人：{config['proofreader']}\t班级__________\t学号_________\t姓名____________")
    meta.alignment=0;meta.paragraph_format.first_line_indent=Pt(72.15)
    meta.paragraph_format.left_indent=Pt(0);meta.paragraph_format.space_after=Pt(5)
    positions=[150.9+extra_comp,245.4+extra_comp+extra_proof,324.15+extra_comp+extra_proof,397.65+extra_comp+extra_proof]
    for pos in positions:meta.paragraph_format.tab_stops.add_tab_stop(Pt(pos))
    for r in meta.runs:style_run(r._r);r.bold=True
    for sec in doc.sections:
        h=sec.header.paragraphs[0]
        h.text=f"{config['cohort']}届{config['grade_subject']}{config['review_stage']}{config['series']}  {config['textbook']}";h.alignment=1
        for r in h.runs:style_run(r._r,size=9)
    body=[e for e in src._element.body if e.tag in [qn('w:p'),qn('w:tbl')]]
    qmap={q['id']:q for q in model['questions']};manifest=[]
    for b in model['blocks']:
        if b['role']=='answer_slot':continue
        if variant=='题目版' and b['role']=='answer':continue
        q=qmap.get(b['qid'])
        if b['role']=='placeholder':
            e=OxmlElement('w:p');p=Paragraph(e,doc._body);p.add_run(f"{q['number']}.")
        else:e=clone_block(body[b['index']],src.part,doc.part)
        if b['role']=='section' and 'display_text' in b:
            patch_text(e,0,len(text(e)),b['display_text'])
        if b['role']=='stem' and q and q.get('source_label'):
            label=q['source_label'];patch_text(e,label['start'],label['end'],'')
        if b['role']=='stem' and q and q['type']=='written':
            m=STEM.match(text(e));patch_text(e,0,m.end(),f"{q['number']}.")
        if b['role']=='stem' and q and q['type']=='choice':
            am=ANSWER.search(text(e))
            if am:patch_text(e,am.start(),am.end(),'')
            sm=STEM.match(text(e));end=sm.end()
            while end<len(text(e)) and text(e)[end].isspace():end+=1
            patch_text(e,0,end,'')
            r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=f"(　\u0020　){q['number']}.";r.append(t)
            e.insert(1 if e.find(qn('w:pPr')) is not None else 0,r)
        format_block(e,doc,b['role'],q is not None and q['type']=='choice',q['number'] if q else 1)
        if b['role']=='placeholder':
            pf=Paragraph(e,doc._body).paragraph_format
            pf.space_after=Pt(168);pf.keep_with_next=False
            pf.page_break_before=False
        if b.get('reserve_page'):Paragraph(e,doc._body).paragraph_format.page_break_before=True
        if b['role']=='stem' and q and q.get('source_label'):
            # Separate prefix, source label and body runs; never change pictures.
            prefix=re.match(r'^\(\s*\)\d+[.]',text(e)) if q['type']=='choice' else STEM.match(text(e))
            end=prefix.end();prefix_text=text(e)[:end];patch_text(e,0,end,'')
            for txt,font in reversed([(prefix_text,'宋体'),('（'+q['source_label']['text']+'）','隶书')]):
                r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=txt;t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');r.append(t);style_run(r,font=font)
                e.insert(1 if e.find(qn('w:pPr')) is not None else 0,r)
        if q and q['type']=='choice' and b['role'] in ['stem','body']:
            content=[x for x in model['blocks'] if x['qid']==q['id'] and x['role'] not in ['answer','answer_slot']]
            for pp in ([e] if e.tag==qn('w:p') else e.xpath('.//w:p')):
                Paragraph(pp,doc._body).paragraph_format.keep_with_next=b is not content[-1]
        doc._element.body.insert(len(doc._element.body)-1,e)
        manifest.append({'qid':b['qid'],'role':b['role'],'kind':e.tag.rsplit('}',1)[-1],'text':norm(text(e)),'images':image_hashes(e,doc.part),'choice':bool(q and q['type']=='choice')})
        if variant=='答案版' and b['role']=='stem' and q and q['type']=='choice':
            # Answer belongs after all options, so emitted just before first analysis below instead.
            pass
        if variant=='题目版' and q and q['type']=='written' and not q.get('placeholder') and b['index']==max(x['index'] for x in model['blocks'] if x['qid']==q['id'] and x['role'] not in ['answer','answer_slot']):
            # Keep answer space on the final content paragraph; an empty final paragraph
            # can become a completely blank extra page in WPS.
            tail=e if e.tag==qn('w:p') else e.xpath('.//w:p')[-1]
            Paragraph(tail,doc._body).paragraph_format.space_after=Pt(config['answer_space_lines']*14)
    # Insert explicit MC answer directly before its first explanation.
    if variant=='答案版':
        offset=2;seen=set()
        for i,b in enumerate(list(manifest)):
            q=qmap.get(b['qid'])
            if q and q['type']=='choice' and b['role']=='answer' and q['id'] not in seen:
                seen.add(q['id']);el=OxmlElement('w:p');p=Paragraph(el,doc._body);p.add_run('【答案】 '+q['answer']);format_block(el,doc,'answer')
                doc._element.body.insert(offset+i,el)
                manifest.insert(i+len(seen)-1,{'qid':q['id'],'role':'answer','kind':'p','text':norm(text(el)),'images':[],'choice':True})
                offset+=1
    # Unique drawing IDs across imported images.
    for i,n in enumerate(doc._element.xpath('//wp:docPr'),1):n.set('id',str(i))
    prune(doc);doc.save(out)
    return manifest
