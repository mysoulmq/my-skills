"""Independent checks against the parsed-source manifest and final DOC roundtrip."""
from pathlib import Path
from collections import Counter
import json,re
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
import pdfplumber
import pypdfium2 as pdfium
from common import norm,text,image_hashes,digest,office,write_json,table_structure


def effective(r,p,doc,prop,attr='val'):
    candidates=[]
    rp=r.find(qn('w:rPr'))
    if rp is not None:
        candidates.append(rp)
        rs=rp.find(qn('w:rStyle'))
        if rs is not None:
            try:candidates.append(doc.styles[rs.get(qn('w:val'))].element.find(qn('w:rPr')))
            except KeyError:pass
    from docx.text.paragraph import Paragraph
    st=Paragraph(p,doc._body).style;seen=set()
    while st is not None and st.style_id not in seen:
        seen.add(st.style_id);candidates.append(st.element.find(qn('w:rPr')));st=st.base_style
    candidates+=doc.styles.element.xpath('./w:docDefaults/w:rPrDefault/w:rPr')
    for c in candidates:
        if c is not None:
            n=c.find(qn('w:'+prop))
            if n is not None and n.get(qn('w:'+attr)) is not None:return n.get(qn('w:'+attr))
    return None

def verify(run):
    run=Path(run);info=json.loads((run/'manifest.json').read_text());errors=[];warnings=[];stats={};actual_common={}
    source_model=json.loads((run/'input.json').read_text())
    from inspect_input import inspect
    from selection import select
    reread=select(inspect(source_model['source']))
    if reread!=source_model:errors.append('筛选模型与原输入重新检查结果不一致')
    source_doc=Document(source_model['source'])
    source_elements=[e for e in source_doc._element.body if e.tag in [qn('w:p'),qn('w:tbl')]]
    paths=[run/'candidates'/f'{v}.doc' for v in ['题目版','答案版']]
    for f in paths:
        if not f.exists() or f.read_bytes()[:8]!=bytes.fromhex('D0CF11E0A1B11AE1'):errors.append(f'{f.name}不是真正的二进制DOC')
    if errors:return {'status':'failed','errors':errors}
    # Always reread final files; never trust earlier inspection artifacts.
    import tempfile
    with tempfile.TemporaryDirectory(prefix='verify-',dir=run) as td:
        roundtrips=office(paths,Path(td)/'docx','docx')
        for v,f in zip(['题目版','答案版'],roundtrips):
            d=Document(f);expected=info['variants'][v]
            allblocks=[e for e in d._element.body if e.tag in [qn('w:p'),qn('w:tbl')]]
            # Template text is outside the question manifest; check it independently.
            groups=[(d.paragraphs[0],['华文楷体','STKaiti'],'28','标题'),(d.paragraphs[1],['宋体','SimSun'],'21','信息栏')]
            groups += [(p,['宋体','SimSun'],'18','页眉') for sec in d.sections for p in sec.header.paragraphs]
            for par,fonts,size,part in groups:
                for r in par._p.xpath('./w:r'):
                    if not text(r).strip():continue
                    if effective(r,par._p,d,'rFonts','eastAsia') not in fonts or effective(r,par._p,d,'sz') != size:errors.append(v+part+'字体或字号偏离样本')

            blocks=[e for e in allblocks[2:] if norm(text(e)) or image_hashes(e,d.part)]
            if len(blocks)!=len(expected):errors.append(f'{v}块数不同: {len(blocks)} != {len(expected)}')
            actual_common[v]=[]
            expected_tables=[table_structure(source_elements[b['index']]) for b in source_model['blocks'] if b.get('kind')=='tbl' and not (v=='题目版' and b['role']=='answer')]
            if [table_structure(e) for e in blocks if e.tag==qn('w:tbl')]!=expected_tables:errors.append(v+'表格行列或合并关系改变')
            for idx,(e,b) in enumerate(zip(blocks,expected)):
                label=f'{v}/{b["qid"] or "section"}/块{idx}'
                if norm(text(e))!=b['text']:errors.append(label+'文字不一致')
                if image_hashes(e,d.part)!=b['images']:errors.append(label+'图片像素不一致或遗失')
                if e.tag.rsplit('}',1)[-1]!=b['kind']:errors.append(label+'表格或段落类型改变')
                if b['role'] not in ['answer','answer_slot']:actual_common[v].append(norm(text(e)))
                for p in [e] if e.tag==qn('w:p') else e.xpath('.//w:p'):
                    cursor=0
                    for r in p.xpath('.//w:r'):
                        run_start=cursor;cursor+=len(text(r))
                        if not norm(text(r)):continue
                        q=next((q for q in source_model['questions'] if q['id']==b['qid']),None)
                        pref=re.match(r'^\(\s*\)\d+[.]',text(p)) if b['choice'] else re.match(r'^\s*\d+[.．、]',text(p))
                        lo=pref.end() if pref else -1
                        hi=lo+len('（'+q['source_label']['text']+'）') if q and q.get('source_label') else -1
                        is_label=bool(b['role']=='stem' and lo<=run_start<hi)
                        if is_label and cursor>hi:errors.append(label+'题源字体范围混入正文')
                        expected_fonts=['隶书','LiSu'] if is_label else ['宋体','SimSun']
                        if effective(r,p,d,'rFonts','eastAsia') not in expected_fonts:errors.append(label+'字体不符合正文或题源规则');break
                        if effective(r,p,d,'sz')!='21':errors.append(label+'非五号');break
                        color=effective(r,p,d,'color')
                        if b['role']=='answer' and color!='FF0000':errors.append(label+'答案解析不是红色');break
                    if b['choice'] and b['role'] in ['stem','body'] and e.tag==qn('w:p'):
                        from docx.text.paragraph import Paragraph
                        pf=Paragraph(p,d._body).paragraph_format
                        # LibreOffice emits start/end attributes, python-docx may not resolve them.
                        inds=p.xpath('./w:pPr/w:ind');ind=inds[0] if inds else None
                        left=int(ind.get(qn('w:left'),ind.get(qn('w:start'),'0'))) if ind is not None else 0
                        hanging=int(ind.get(qn('w:hanging'),'0')) if ind is not None else 0
                        number=next(q['number'] for q in source_model['questions'] if q['id']==b['qid'])
                        expected_left=840+105*len(str(number))
                        if abs(left-expected_left)>2:errors.append(label+f'正文左边界不正确:{left}')
                        if b['role']=='stem' and (abs(hanging-expected_left)>2 or not re.match(r'^\(\s*\)\d+[.]\S',text(e))):errors.append(label+'题号后存在空白或悬挂不正确')
            for e,b in zip(blocks,expected):
                if b['role']=='placeholder':
                    from docx.text.paragraph import Paragraph
                    pf=Paragraph(e,d._body).paragraph_format
                    if pf.space_after is None or pf.space_after.pt<160:errors.append(v+'综合题空位不足')
            with ZipFile(f) as z:
                for name in z.namelist():
                    if name.endswith('.xml') and name.startswith('word/'):
                        data=z.read(name)
                        if any(x in data for x in [b'<w:del ',b'<w:ins ',b'<w:comment ',b'<w:vanish/>',b'<w:vanish w:val="true"']):errors.append(v+'包含隐藏/修订/批注')
            # Metadata check verifies actual saved default or overridden names.
            meta=text(allblocks[1]);cfg=info['config']
            if cfg['compiler'] not in meta or cfg['proofreader'] not in meta:errors.append(v+'编制/校对信息错误')
            meta_runs=allblocks[1].xpath('.//w:r[w:t]')
            if any(not r.xpath('./w:rPr/w:b[not(@w:val) or @w:val="1" or @w:val="true"]') for r in meta_runs):errors.append(v+'信息栏未按样本加粗')
            header=''.join(p.text for s in d.sections for p in s.header.paragraphs)
            expected_header=f"{cfg['cohort']}届{cfg['grade_subject']}{cfg['review_stage']}{cfg['series']}  {cfg['textbook']}"
            if norm(expected_header) not in norm(header):errors.append(v+'页眉信息错误')
            stats[v]={'blocks':len(blocks),'sha256':digest(run/'candidates'/f'{v}.doc')}
    if actual_common['题目版']!=actual_common['答案版']:errors.append('两版题目内容不一致')
    # Compare independently with the original source, not just the generator's manifest.
    model=json.loads((run/'input.json').read_text())
    if digest(model['source'])!=model['source_sha256']:errors.append('输入文件在转换后发生改变')
    for v in ['题目版','答案版']:
        generated=info['variants'][v]
        for q in model['questions']:
            if q.get('placeholder'):
                bs=[b for b in generated if b['qid']==q['id']]
                if len(bs)!=1 or bs[0]['text']!=f"{q['source_number']}." or bs[0]['images']:errors.append(v+'/'+q['id']+'综合题空位残留内容或原编号改变')
                continue
            source_blocks=[b for b in model['blocks'] if b['qid']==q['id'] and b['role'] not in ['answer','answer_slot']]
            source_text=''.join(b['text'] for b in source_blocks)
            from inspect_input import ANSWER,STEM
            stem=source_blocks[0]['text'];am=ANSWER.search(stem) if q['type']=='choice' else None
            if am:stem=stem[:am.start()]+stem[am.end():]
            if q.get('source_label'):
                tag=q['source_label'];stem=stem[:tag['start']]+'（'+tag['text']+'）'+stem[tag['end']:]
            stem=STEM.sub('',stem,count=1)
            source_text=stem+''.join(b['text'] for b in source_blocks[1:])
            actual=''.join(b['text'] for b in generated if b['qid']==q['id'] and b['role'] not in ['answer','answer_slot'])
            actual=re.sub(r'^\(\)\d+[.]','',actual) if q['type']=='choice' else STEM.sub('',actual,count=1)
            if norm(source_text)!=norm(actual):errors.append(v+'/'+q['id']+'与原始题目内容不一致')
            original_images=[h for b in source_blocks for h in b.get('images',[])]
            output_images=[h for b in generated if b['qid']==q['id'] and b['role'] not in ['answer','answer_slot'] for h in b['images']]
            if original_images!=output_images:errors.append(v+'/'+q['id']+'题目图片不完整')
            if v=='答案版':
                original_answer=norm(''.join(b['text'] for b in model['blocks'] if b['qid']==q['id'] and b['role']=='answer'))
                rendered_answer=''.join(b['text'] for b in generated if b['qid']==q['id'] and b['role']=='answer')
                if q['type']=='choice':rendered_answer=rendered_answer.removeprefix('【答案】'+q['answer'])
                if original_answer!=rendered_answer:errors.append(v+'/'+q['id']+'原解析内容改变')
    pdfs=office(paths,run/'qa/pdf','pdf')
    for v,pdf in zip(['题目版','答案版'],pdfs):
        raster=pdfium.PdfDocument(pdf)
        with pdfplumber.open(pdf) as d:
            from layout_audit import audit_question_layout
            layout=audit_question_layout(d,source_model,info['variants'][v])
            errors.extend(v+x for x in layout['errors']);stats[v]['layout']=layout
            stats[v]['pages']=len(d.pages);stats[v]['pdf_sha256']=digest(pdf)
            title_chars=[c for c in d.pages[0].chars if re.search('[\u4e00-\u9fff]',c['text']) and abs(c['size']-14)<.2]
            if title_chars and any('STKaiti' not in c['fontname'] for c in title_chars):warnings.append(v+'标题指定华文楷体，本机预览使用替代字体；真实字体环境的字形和换行未验证')
            label_seen=[];label_positions=set()
            expected_labels={q['number']:norm('（'+q['source_label']['text']+'）') for q in source_model['questions'] if q.get('source_label') and not q.get('placeholder')}
            for page in d.pages:
                for line in page.extract_text_lines():
                    cs=[c for c in line['chars'] if c['text'].strip()];visible=''.join(c['text'] for c in cs)
                    match=re.match(r'^(?:\(\))?(\d+)\.(（[^）]+）)',visible)
                    if not match:continue
                    num=int(match[1]);label_seen.append(num)
                    if expected_labels.get(num)!=match[2]:errors.append(v+f'/第{num}题题源括号或内容不正确')
                    start=match.start(2);end=match.end(2)
                    label_positions.update((c['page_number'],c['x0'],c['top']) for c in cs[start:end])
                    if any(abs(c['size']-10.5)>.2 for c in cs[start:end]):errors.append(v+f'/第{num}题题源不是五号')
                    if any('LiSu' not in c['fontname'] for c in cs[start:end]):warnings.append(v+f'/第{num}题文档指定隶书，本机预览使用替代字体；真实隶书环境的字形和换行未验证')
            if sorted(label_seen)!=sorted(expected_labels):errors.append(v+'题源标注缺失或重复')

            meta_lines=[norm(line) for line in (d.pages[0].extract_text() or '').splitlines() if '编制人' in line]
            if not any(all(key in line for key in ['编制人','校对人','班级__________','学号_________','姓名____________']) for line in meta_lines):errors.append(v+'学生信息栏换行或内容不完整')
            # Verify actual visible start against the reference, not just no-wrap.
            words=d.pages[0].extract_words()
            starts=[w for w in words if w['text'].startswith('编制人')]
            if len(starts)!=1 or abs(starts[0]['x0']-117.5)>2:errors.append(v+'信息栏起点偏离样本117.5pt')
            expected_x={'编制人':117.5,'校对人':196.25+max(0,len(info['config']['compiler'])-3)*10.5}
            shift=max(0,len(info['config']['compiler'])-3)*10.5+max(0,len(info['config']['proofreader'])-3)*10.5
            expected_x.update({'班级':290.75+shift,'学号':369.5+shift,'姓名':443+shift})
            for label,x in expected_x.items():
                found=[w for w in words if w['text'].startswith(label)]
                if len(found)!=1 or abs(found[0]['x0']-x)>2:errors.append(v+'信息栏'+label+'字段位置偏离样本')
            if starts:
                linechars=[c for c in d.pages[0].chars if abs(c['top']-starts[0]['top'])<2 and c['text'].strip()]
                if any(c['x1']>d.pages[0].width-45.35+1 for c in linechars):errors.append(v+'信息栏超出右页边距')
            target=run/'qa/pages'/v;target.mkdir(parents=True,exist_ok=True)
            for i,page in enumerate(d.pages):
                raster[i].render(scale=1.4).to_pil().save(target/f'page-{i+1}.png')
                chars=[c for c in page.chars if c['text'].strip()]
                for c in chars:
                    if (c['page_number'],c['x0'],c['top']) not in label_positions and re.search('[\u4e00-\u9fff]',c['text']) and abs(c['size']-10.5)<0.2 and 'SimSun' not in c['fontname']:
                        errors.append(f'{v}第{i+1}页正文渲染字体发生替换');break
                if not chars:errors.append(f'{v}第{i+1}页为空页')
                if any(c['x0']<0 or c['top']<0 or c['x1']>page.width+1 or c['bottom']>page.height+1 for c in chars):errors.append(f'{v}第{i+1}页文字超出页面')
        raster.close()
    deletion_info=None
    if model.get('removed_questions'):
        report=run/'candidates/删除记录.pdf'
        if not report.exists():errors.append('缺少删除记录PDF')
        else:
            with pdfplumber.open(report) as pdf:
                txt=norm(''.join(p.extract_text() or '' for p in pdf.pages))
                if any(norm(r['source_label']) not in txt for r in model['removed_questions']):errors.append('删除PDF缺少题源记录')
                deletion_info={'sha256':digest(report),'pages':len(pdf.pages)}
            pdf=pdfium.PdfDocument(report);target=run/'qa/pages/删除记录';target.mkdir(parents=True,exist_ok=True)
            for i,page in enumerate(pdf):page.render(scale=1.4).to_pil().save(target/f'page-{i+1}.png')
            pdf.close()
    result={'status':'failed' if errors else 'awaiting_visual_review','errors':sorted(set(errors)),'warnings':sorted(set(warnings)),'files':stats,'review_issues':info['review_issues'],'note':'自动检查不替代逐页视觉及教研审核。'}
    if deletion_info:result['deletion_report']=deletion_info
    write_json(run/'verification.json',result);return result
