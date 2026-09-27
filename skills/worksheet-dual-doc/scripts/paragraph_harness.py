#!/usr/bin/env python3
"""Small offline saved-DOC regression: good cases + deliberately broken files.

Run with the same bundled Python as worksheet.py. No model or private fixtures.
"""
import json,tempfile
from pathlib import Path
from docx import Document
from docx.shared import Pt
import pdfplumber
from common import ROOT,office,text,norm
from inspect_input import inspect
from generate import generate
from paragraph_rules import audit_paragraphs
from verify_output import effective
from layout_audit import audit_question_layout

def run():
    results=[]
    with tempfile.TemporaryDirectory(prefix='worksheet-paragraph-harness-') as td:
        root=Path(td);src=root/'input.docx';d=Document()
        d.add_paragraph('高效作业27 学习借鉴外来文化的有益成果')
        d.add_paragraph('一、选择题')
        d.add_paragraph('1.2025·江苏高考'+('文化交流推动社会发展。'*16)+'\t( B )\u3000\u3000')
        d.add_paragraph('①文化交流 ②文化发展 ③文化传承 ④文化创新')
        d.add_paragraph('A.①②\tB.①③\tC.②④\tD.③④')
        d.add_paragraph('【解析】故选B。')
        d.add_paragraph('二、综合题')
        d.add_paragraph('2.阅读材料，回答问题。')
        d.add_paragraph('材料一：文化交流推动不同文明相互借鉴。')
        prompt='结合材料，运用文化的知识，说明交流的意义。（9分）'
        d.add_paragraph(prompt)
        d.add_paragraph('【答案】①促进交流。②促进发展。')
        d.save(src);model=inspect(src);assert not model['errors'],model['errors']
        cfg=json.loads((ROOT/'assets/defaults.json').read_text())
        for variant in ('题目版','答案版'):
            internal=root/(variant+'.docx');manifest=generate(src,model,cfg,variant,internal)
            def check(name):
                saved=office([internal],root/name,'doc:MS Word 97')[0]
                back=office([saved],root/(name+'-back'),'docx')[0]
                final=Document(back);blocks=list(final._element.body)[2:-1]
                errors=audit_paragraphs(final,blocks,model,effective)
                pdf=office([saved],root/(name+'-pdf'),'pdf')[0]
                with pdfplumber.open(pdf) as pages:geometry=audit_question_layout(pages,model,manifest)
                return final,errors,geometry['errors']
            good,errors,geometry=check(variant+'-good')
            assert not errors+geometry,errors+geometry
            assert any(norm(p.text)==norm(prompt) for p in good.paragraphs)
            material=next(p for p in good.paragraphs if p.text.startswith('材料一'))
            assert not any(r.bold for r in material.runs),'材料被误加粗'
            results.append(variant+': clean saved DOC passed')
            # Mutations keep meaningful text identical: the harness must detect layout errors.
            bad=Document(internal)
            p=next(p for p in bad.paragraphs if p.text==prompt)
            for r in p.runs:r.bold=False
            p=next(p for p in bad.paragraphs if p.text.startswith('('))
            p.add_run('\t');p.paragraph_format.line_spacing=Pt(24)
            bad.save(internal)
            _,errors,geometry=check(variant+'-bad')
            for expected in ['设问未全部加粗','末尾残留空白','行距未固定']:
                assert any(expected in e for e in errors),(expected,errors)
            assert any('题干行距异常' in e for e in geometry),geometry
            results.append(variant+': missing bold / trailing tab / inflated line height rejected')
    return {'status':'passed','cases':results}

if __name__=='__main__':print(json.dumps(run(),ensure_ascii=False,indent=2))
