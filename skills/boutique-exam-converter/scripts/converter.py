#!/usr/bin/env python3
"""Convert a paired publisher exam into three school worksheet DOCX files."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from PIL import Image
from pypdf import PdfReader
import pdfplumber


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
QUESTION = re.compile(r"^\s*(\d+)[.．、]\s*")
CHOICE_ANSWER = re.compile(r"【答案】\s*([A-DＡ-Ｄ](?:\s*[,，、]?\s*[A-DＡ-Ｄ])*)")
WRITTEN_SECTION = re.compile(r"(综合题|主观题|非选择题)")
CHOICE_SECTION = re.compile(r"选择题")
IMAGE_WORDS = re.compile(r"漫画|下图|图示|图中|曲线|柱状图|表格")


def norm(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path: Path, value) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def block_text(element) -> str:
    values = []
    for node in element.iter():
        if node.getparent() is None or node.getparent().tag != qn("w:r"):
            continue
        if node.tag == qn("w:t"):
            values.append(node.text or "")
        elif node.tag == qn("w:tab"):
            values.append("\t")
        elif node.tag in (qn("w:br"), qn("w:cr")):
            values.append("\n")
    return "".join(values)


def image_ids(element):
    return element.xpath(".//a:blip/@r:embed")


def image_hash(blob: bytes) -> str:
    image = Image.open(BytesIO(blob)).convert("RGBA")
    return hashlib.sha256(str(image.size).encode() + image.tobytes()).hexdigest()


def block_image_hashes(element, part):
    result = []
    for rid in image_ids(element):
        if rid in part.related_parts:
            result.append(image_hash(part.related_parts[rid].blob))
    return result


def body_blocks(doc: Document):
    return [
        element
        for element in doc._element.body
        if element.tag in (qn("w:p"), qn("w:tbl"))
    ]


def section_kind(text: str):
    clean = text.strip()
    if not clean or "题" not in clean:
        return None
    # Section headings are short structural lines. Explanations such as
    # “本题为主观题” must remain inside the current question.
    if not (re.match(r"^[一二三四五六七八九十]+[、．.]", clean) or clean in ("非选择题 部分", "非选择题部分")):
        return None
    if WRITTEN_SECTION.search(clean):
        return "written"
    if CHOICE_SECTION.search(clean):
        return "choice"
    return None


def parse_questions(path: Path):
    doc = Document(path)
    questions = []
    headings = {"choice": [], "written": []}
    current = None
    kind = None
    for index, element in enumerate(body_blocks(doc)):
        text = block_text(element)
        detected = section_kind(text)
        if detected:
            kind = detected
            headings[kind].append(text.strip())
            current = None
            continue
        match = QUESTION.match(text) if element.tag == qn("w:p") else None
        if match and kind is not None:
            current = {
                "number": int(match.group(1)),
                "kind": kind or "unknown",
                "blocks": [],
            }
            questions.append(current)
        if current is not None:
            if text.strip() or block_image_hashes(element, doc.part) or element.tag == qn("w:tbl"):
                current["blocks"].append(element)

    # Publisher DOCX files sometimes put a question image immediately before its stem.
    for pos in range(1, len(questions)):
        previous = questions[pos - 1]
        current = questions[pos]
        if not current["blocks"] or not IMAGE_WORDS.search(block_text(current["blocks"][0])):
            continue
        moved = []
        while previous["blocks"]:
            tail = previous["blocks"][-1]
            if block_text(tail).strip() or not block_image_hashes(tail, doc.part):
                break
            moved.insert(0, previous["blocks"].pop())
        if moved:
            current["blocks"][1:1] = moved
    return doc, questions, headings


def visible_question_prefix(question):
    text_parts = []
    image_parts = []
    for block in question["blocks"]:
        text = block_text(block)
        if "【答案】" in text:
            before = text.split("【答案】", 1)[0]
            if before:
                text_parts.append(before)
            break
        text_parts.append(text)
        image_parts.extend(block_image_hashes(block, question["doc"].part))
    return norm("".join(text_parts)), image_parts


def split_answer(question):
    texts = [block_text(block).strip() for block in question["blocks"]]
    answer_at = next((i for i, text in enumerate(texts) if "【答案】" in text), None)
    if answer_at is None:
        return None
    parse_at = next((i for i in range(answer_at, len(texts)) if texts[i] == "【解析】"), None)
    if parse_at is None:
        parse_at = next((i for i in range(answer_at, len(texts)) if texts[i].startswith("【解析】")), len(texts))
    answers = [text for text in texts[answer_at:parse_at] if text]
    details = []
    for text in texts[parse_at:]:
        if not text or text == "【解析】":
            continue
        if text.startswith("【解析】"):
            text = text[len("【解析】") :]
        if text:
            details.append(text)
    return {"answer": answers, "details": details}


def inspect_pair(question_path: Path, analysis_path: Path):
    question_doc, original, headings = parse_questions(question_path)
    analysis_doc, analysed, _ = parse_questions(analysis_path)
    for question in original:
        question["doc"] = question_doc
    for question in analysed:
        question["doc"] = analysis_doc
    errors = []
    review_issues = []
    if not original:
        errors.append("原卷版没有识别出题目")
    if not analysed:
        errors.append("解析版没有识别出题目")
    original_numbers = [q["number"] for q in original]
    analysed_numbers = [q["number"] for q in analysed]
    if original_numbers != analysed_numbers:
        errors.append(f"两份输入题号不一致：原卷{original_numbers}；解析版{analysed_numbers}")
    if len(set(original_numbers)) != len(original_numbers):
        errors.append("原卷版存在重复题号")

    rows = []
    by_number = {q["number"]: q for q in analysed}
    for original_q in original:
        number = original_q["number"]
        analysis_q = by_number.get(number)
        row = {"number": number, "kind": original_q["kind"]}
        if analysis_q is None:
            row["status"] = "missing"
            rows.append(row)
            continue
        left_text, left_images = visible_question_prefix(original_q)
        right_text, right_images = visible_question_prefix(analysis_q)
        if left_text != right_text:
            errors.append(f"第{number}题题目文字不一致")
            row["question_text_match"] = False
            row["original_text"] = left_text
            row["analysis_text"] = right_text
        else:
            row["question_text_match"] = True
        if left_images != right_images:
            errors.append(f"第{number}题题目图片不一致")
            row["question_images_match"] = False
        else:
            row["question_images_match"] = True
        answer = split_answer(analysis_q)
        if answer is None:
            errors.append(f"第{number}题解析版缺少【答案】")
        else:
            row["answer_paragraphs"] = len(answer["answer"])
            row["detail_paragraphs"] = len(answer["details"])
            if original_q["kind"] == "choice":
                answer_text = "".join(answer["answer"])
                choice = CHOICE_ANSWER.search(answer_text)
                if not choice:
                    errors.append(f"第{number}题缺少可识别的选择题答案字母")
                else:
                    letters = choice.group(1).translate(str.maketrans("ＡＢＣＤ", "ABCD"))
                    row["choice_answer"] = re.sub(r"[^A-D]", "", letters)
                    conclusions = re.findall(r"(?:故选|故答案选|故答案为)\s*([A-D])", "".join(answer["details"]))
                    if conclusions and any(value not in row["choice_answer"] for value in conclusions):
                        review_issues.append(
                            f"第{number}题答案{row['choice_answer']}与详解结论{conclusions}疑似冲突，需教研确认"
                        )
        row["status"] = "matched" if row.get("question_text_match") and row.get("question_images_match") else "mismatch"
        rows.append(row)

    serializable_questions = []
    for question in original:
        serializable_questions.append(
            {
                "number": question["number"],
                "kind": question["kind"],
                "text": norm("".join(block_text(block) for block in question["blocks"])),
                "images": [
                    value
                    for block in question["blocks"]
                    for value in block_image_hashes(block, question_doc.part)
                ],
            }
        )
    return {
        "question_source": str(question_path.resolve()),
        "analysis_source": str(analysis_path.resolve()),
        "question_sha256": digest(question_path),
        "analysis_sha256": digest(analysis_path),
        "question_count": len(original),
        "choice_count": sum(q["kind"] == "choice" for q in original),
        "written_count": sum(q["kind"] == "written" for q in original),
        "questions": serializable_questions,
        "pair_review": rows,
        "headings": headings,
        "errors": errors,
        "review_issues": review_issues,
    }, question_doc, original, analysis_doc, analysed


def set_run_font(run, name="宋体", size=10.5, color=None, bold=None):
    run.font.name = name
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def copy_ppr(target, prototype):
    old = target.find(qn("w:pPr"))
    if old is not None:
        target.remove(old)
    source = prototype.find(qn("w:pPr"))
    if source is not None:
        target.insert(0, deepcopy(source))


def clone_block(element, source_part, dest_part):
    node = deepcopy(element)
    for blip in node.xpath(".//a:blip"):
        rid = blip.get(qn("r:embed"))
        if rid and rid in source_part.related_parts:
            new_rid, _ = dest_part.get_or_add_image(BytesIO(source_part.related_parts[rid].blob))
            blip.set(qn("r:embed"), new_rid)
    for link in list(node.xpath(".//w:hyperlink")):
        parent = link.getparent()
        position = parent.index(link)
        for child in list(link):
            parent.insert(position, child)
            position += 1
        parent.remove(link)
    return node


def patch_visible_text(element, start, end, replacement=""):
    atoms = [
        node
        for node in element.iter()
        if node.tag in (qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr"))
        and node.getparent().tag == qn("w:r")
    ]
    position = 0
    inserted = False
    for node in atoms:
        value = node.text or "" if node.tag == qn("w:t") else "\t" if node.tag == qn("w:tab") else "\n"
        left, right = position, position + len(value)
        position = right
        if right <= start or left >= end:
            continue
        new_value = value[: max(0, start - left)]
        if not inserted:
            new_value += replacement
            inserted = True
        new_value += value[min(len(value), end - left) :]
        node.tag = qn("w:t")
        node.text = new_value
        node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def add_prefix(element, value):
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = value
    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    run.append(text)
    element.insert(1 if element.find(qn("w:pPr")) is not None else 0, run)


def style_element(element, font, prototype_ppr=None, size=10.5):
    paragraphs = [element] if element.tag == qn("w:p") else element.xpath(".//w:p")
    for paragraph in paragraphs:
        if prototype_ppr is not None:
            copy_ppr(paragraph, prototype_ppr)
        for run_element in paragraph.xpath(".//w:r"):
            from docx.text.run import Run

            run = Run(run_element, None)
            set_run_font(run, font, size)


def prune(doc):
    used = set(image_ids(doc._element))
    for rid, relation in list(doc.part.rels.items()):
        kind = relation.reltype.rsplit("/", 1)[-1]
        if kind == "image" and rid not in used:
            doc.part.drop_rel(rid)
    for index, node in enumerate(doc._element.xpath("//wp:docPr"), 1):
        node.set("id", str(index))
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""


def nonempty_paragraphs(doc):
    return [paragraph for paragraph in doc.paragraphs if paragraph.text.strip()]


def template_labels(kind, config):
    title = f"{config['school']}{config['cohort']}届{config['subject']}{config['stage']}{config['document_kind']}"
    if kind == "analysis":
        return title, f"{config['practice_series']}{config['practice_number']}解析", None, None
    suffix = " 选择题" if kind == "selection" else "主观题"
    subtitle = f"{config['practice_series']}{config['practice_number']}{suffix}"
    meta = (
        f"编制人：{config['compiler']}      校对人：{config['proofreader']}          "
        f"{config['class_label']}             {config['student_number_label']}          "
        f"{config['student_name_label']}              "
    )
    section = "一．选择题" if kind == "selection" else "二．主观题"
    return title, subtitle, meta, section


def patch_paragraph_preserving_objects(paragraph, value):
    old = block_text(paragraph._p)
    if old == value:
        return
    patch_visible_text(paragraph._p, 0, len(old), value)


def customize_reference(kind, config, output):
    """Edit labels in-place without rebuilding paragraphs or deleting anchored artwork."""
    doc = Document(ASSETS / f"reference-{kind}.docx")
    anchors = nonempty_paragraphs(doc)
    title, subtitle, meta, section = template_labels(kind, config)
    patch_paragraph_preserving_objects(anchors[0], title)
    patch_paragraph_preserving_objects(anchors[1], subtitle)
    if kind != "analysis":
        patch_paragraph_preserving_objects(anchors[2], meta)
        patch_paragraph_preserving_objects(anchors[3], section)
    prune(doc)
    doc.save(output)


def prepare_template(kind, config):
    template = ASSETS / f"reference-{kind}.docx"
    doc = Document(template)
    blocks = body_blocks(doc)
    if kind == "selection":
        start = next(i for i, block in enumerate(blocks) if re.match(r"^\s*（\s*）\d+[.]", block_text(block)))
        prototype_stem = deepcopy(blocks[start])
        prototype_body = deepcopy(blocks[start + 1])
    elif kind == "written":
        start = next(i for i, block in enumerate(blocks) if QUESTION.match(block_text(block)))
        prototype_stem = deepcopy(blocks[start])
        prototype_body = deepcopy(blocks[start + 2])
    else:
        start = next(i for i, block in enumerate(blocks) if re.match(r"^\s*1[.]【答案】", block_text(block)))
        prototype_stem = deepcopy(blocks[start])
        prototype_body = deepcopy(blocks[start])
    for block in blocks[start:]:
        doc._element.body.remove(block)

    anchors = nonempty_paragraphs(doc)
    title, subtitle, meta, section = template_labels(kind, config)
    if kind == "analysis":
        patch_paragraph_preserving_objects(anchors[0], title)
        patch_paragraph_preserving_objects(anchors[1], subtitle)
    else:
        patch_paragraph_preserving_objects(anchors[0], title)
        patch_paragraph_preserving_objects(anchors[1], subtitle)
        patch_paragraph_preserving_objects(anchors[2], meta)
        patch_paragraph_preserving_objects(anchors[3], section)
    return doc, prototype_stem, prototype_body


def insert_before_section(doc, element):
    doc._element.body.insert(len(doc._element.body) - 1, element)


def strip_paragraph_numbering(element):
    """Remove inherited Word list numbering when the source already carries visible labels."""
    if element.tag != qn("w:p"):
        return
    for num_pr in element.xpath("./w:pPr/w:numPr"):
        num_pr.getparent().remove(num_pr)


def make_question_doc(kind, config, source_doc, questions, output):
    doc, stem_proto, body_proto = prepare_template(kind, config)
    stem_ppr = stem_proto
    body_ppr = body_proto
    selected = [q for q in questions if q["kind"] == ("choice" if kind == "selection" else "written")]
    for question in selected:
        for position, source_block in enumerate(question["blocks"]):
            element = clone_block(source_block, source_doc.part, doc.part)
            text = block_text(element)
            is_stem = position == 0 and element.tag == qn("w:p")
            if kind == "selection" and is_stem:
                match = QUESTION.match(text)
                if match:
                    patch_visible_text(element, 0, match.end(), "")
                    current = block_text(element)
                    trailing = re.search(r"[（(]\s*[)）]\s*$", current)
                    if trailing:
                        patch_visible_text(element, trailing.start(), trailing.end(), "")
                    add_prefix(element, f"（    ）{question['number']}. ")
                style_element(element, "宋体", stem_ppr)
            elif kind == "selection":
                style_element(element, "宋体", body_ppr)
            else:
                material = is_stem or bool(re.match(r"^\s*材料[一二三四五六七八九十\d]*", text))
                style_element(element, "楷体" if material else "宋体", stem_ppr if material else body_ppr)
                # style_element copies the sample paragraph properties, including its
                # list definition. The source already contains every desired visible
                # subquestion label, so inherited numbering would create duplicates.
                strip_paragraph_numbering(element)
            insert_before_section(doc, element)
    prune(doc)
    doc.save(output)
    return [q["number"] for q in selected]


def answer_chunks(answer_parts):
    chunks = []
    current = ""
    for part in answer_parts:
        if re.fullmatch(r"[（(][二三四五六七八九十2-9]+[)）]", part) and current:
            chunks.append(current)
            current = part
        else:
            current += part
    if current:
        chunks.append(current)
    return chunks


def add_analysis_paragraph(doc, prototype, text, *, bold=False):
    paragraph = doc.add_paragraph()
    copy_ppr(paragraph._p, prototype)
    pieces = re.split(r"(【答案】)", text)
    for piece in pieces:
        if not piece:
            continue
        run = paragraph.add_run(piece)
        set_run_font(run, "宋体", 10.5, "2E75B6" if piece == "【答案】" else "000000", bold if piece != "【答案】" else False)
    return paragraph


def analysis_records(analysis_questions):
    records = []
    for question in analysis_questions:
        split = split_answer(question)
        if split is None:
            continue
        records.append({"number": question["number"], "kind": question["kind"], **split})
    return records


def generate_analysis(config, records, written_heading, detail_counts, output):
    doc, prototype, _ = prepare_template("analysis", config)
    for record in records:
        if record["kind"] != "choice":
            continue
        content = "".join(record["answer"] + record["details"])
        add_analysis_paragraph(doc, prototype, f"{record['number']}.{content}")
    add_analysis_paragraph(doc, prototype, "非选择题 部分", bold=True)
    heading = written_heading or f"三、综合题（本大题共{sum(r['kind'] == 'written' for r in records)}小题）"
    add_analysis_paragraph(doc, prototype, heading, bold=True)
    for record in records:
        if record["kind"] != "written":
            continue
        chunks = answer_chunks(record["answer"])
        if not chunks:
            chunks = ["【答案】"]
        count = detail_counts.get(record["number"], 0)
        details = record["details"][:count]
        if details:
            chunks[-1] += "".join(details)
        for index, chunk in enumerate(chunks):
            prefix = f"{record['number']}." if index == 0 else ""
            add_analysis_paragraph(doc, prototype, prefix + chunk)
    prune(doc)
    doc.save(output)


def runtime_bin(name):
    python = Path(sys.executable).resolve()
    candidates = []
    for parent in python.parents:
        candidates.extend([parent / "bin" / "override" / name, parent / "bin" / "fallback" / name])
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    found = shutil.which(name)
    if found:
        return found
    raise RuntimeError(f"缺少运行工具{name}；请先加载workspace dependencies")


def render_pdf(docx: Path, output_dir: Path):
    output_dir.mkdir(parents=True, exist_ok=True)
    pdf = output_dir / f"{docx.stem}.pdf"
    if pdf.exists():
        pdf.unlink()
    temp_parent = "/private/tmp" if Path("/private/tmp").exists() else None
    with tempfile.TemporaryDirectory(prefix="boutique-lo-", dir=temp_parent) as profile:
        env = dict(os.environ)
        from fonts import resolve_fonts

        font_info = resolve_fonts()
        font_dir = Path(font_info["directory"])
        fontconfig = Path(profile) / "fonts.conf"
        directories = [
            font_dir,
            Path("/System/Library/Fonts"),
            Path("/Library/Fonts"),
            Path.home() / "Library/Fonts",
            Path("/usr/share/fonts"),
        ]
        fontconfig.write_text(
            "<fontconfig>"
            + "".join(f"<dir>{escape(str(path))}</dir>" for path in directories if path.exists())
            + f"<cachedir>{escape(str(Path(profile) / 'fontcache'))}</cachedir>"
            + "<alias><family>宋体</family><prefer><family>SimSun</family></prefer></alias>"
            + "<alias><family>楷体</family><prefer><family>KaiTi</family><family>SimSun</family></prefer></alias>"
            + "</fontconfig>",
            encoding="utf-8",
        )
        env["FONTCONFIG_FILE"] = str(fontconfig)
        command = [
            runtime_bin("soffice"),
            f"-env:UserInstallation={Path(profile).as_uri()}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(docx.resolve()),
        ]
        result = subprocess.run(command, capture_output=True, text=True, timeout=180, env=env)
        if result.returncode or not pdf.exists() or not pdf.stat().st_size:
            raise RuntimeError(f"DOCX渲染失败：{result.stdout}\n{result.stderr}")
    return pdf, len(PdfReader(str(pdf)).pages)


def last_page_fill(pdf: Path):
    """Approximate body fill, excluding the header and footer/page-number bands."""
    with pdfplumber.open(pdf) as document:
        page = document.pages[-1]
        usable_bottom = page.height - 120
        chars = [
            char
            for char in page.chars
            if char.get("text", "").strip() and char["top"] > 35 and char["bottom"] < page.height - 100
        ]
        if not chars:
            return 0.0
        return round(min(1.0, max(char["bottom"] for char in chars) / usable_bottom), 3)


def render_pngs(pdf: Path, output_dir: Path):
    output_dir.mkdir(parents=True, exist_ok=True)
    prefix = output_dir / "page"
    result = subprocess.run(
        [runtime_bin("pdftoppm"), "-png", "-r", "144", str(pdf), str(prefix)],
        capture_output=True,
        text=True,
        timeout=180,
    )
    if result.returncode:
        raise RuntimeError(f"PDF转PNG失败：{result.stderr}")
    pages = sorted(output_dir.glob("page-*.png"))
    if not pages:
        raise RuntimeError("PDF转PNG后没有页面")
    return pages


def choose_detail_fill(config, records, written_heading, work_dir):
    target = int(config["target_analysis_pages"])
    selected = {record["number"]: 0 for record in records if record["kind"] == "written"}
    trials = []

    def trial(label):
        docx = work_dir / f"analysis-{len(trials):03d}.docx"
        generate_analysis(config, records, written_heading, selected, docx)
        pdf, pages = render_pdf(docx, work_dir / "pdf")
        entry = {
            "label": label,
            "docx": str(docx),
            "pdf": str(pdf),
            "pages": pages,
            "last_page_fill": last_page_fill(pdf),
            "detail_counts": dict(selected),
        }
        trials.append(entry)
        return entry

    best = trial("no-written-details")
    if best["pages"] > target:
        return best, trials, [f"仅参考答案已超过目标{target}页，未恢复主观题详解"]

    threshold = float(config["analysis_fill_threshold"])
    if best["pages"] == target and best["last_page_fill"] >= threshold:
        return best, trials, []
    written = [record for record in records if record["kind"] == "written" and record["details"]]
    stop = False
    for record in reversed(written):
        number = record["number"]
        full = len(record["details"])
        selected[number] = full
        candidate = trial(f"q{number}-full")
        if candidate["pages"] <= target:
            best = candidate
            if candidate["pages"] == target and candidate["last_page_fill"] >= threshold:
                stop = True
                break
            continue
        selected[number] = 0
        for count in range(1, full + 1):
            selected[number] = count
            candidate = trial(f"q{number}-prefix-{count}")
            if candidate["pages"] <= target:
                best = candidate
                if candidate["pages"] == target and candidate["last_page_fill"] >= threshold:
                    stop = True
                    break
            else:
                selected[number] = count - 1
                stop = True
                break
        if stop:
            break
    warnings = []
    if best["pages"] < target:
        warnings.append(f"恢复全部可用主观题详解后仍只有{best['pages']}页，未使用空行撑页")
    return best, trials, warnings


def parse_detail_counts(items, records):
    available = {
        record["number"]: len(record["details"])
        for record in records
        if record["kind"] == "written"
    }
    selected = {number: 0 for number in available}
    for item in items:
        if "=" not in item:
            raise ValueError("详解保留量必须使用题号=段落数，例如 --detail-count 27=4")
        number_text, count_text = item.split("=", 1)
        try:
            number = int(number_text)
            count = int(count_text)
        except ValueError as exc:
            raise ValueError(f"无效详解保留量：{item}") from exc
        if number not in available:
            raise ValueError(f"第{number}题不是可配置的主观题")
        if count < 0 or count > available[number]:
            raise ValueError(f"第{number}题详解段落数应在0至{available[number]}之间")
        selected[number] = count
    return selected


def make_manual_detail_choice(config, records, written_heading, work_dir, items):
    selected = parse_detail_counts(items, records)
    docx = work_dir / "analysis-wps-choice.docx"
    generate_analysis(config, records, written_heading, selected, docx)
    pdf, pages = render_pdf(docx, work_dir / "pdf")
    entry = {
        "label": "manual-wps-selection",
        "docx": str(docx),
        "pdf": str(pdf),
        "pages": pages,
        "last_page_fill": last_page_fill(pdf),
        "detail_counts": selected,
    }
    return entry, [entry], ["详解范围由--detail-count指定；清单中的PDF页数仅供辅助，最终以WPS实测为准"]


def output_names(config):
    prefix = f"{config['stage']} {config['practice_series']}{config['practice_number']}"
    return {
        "selection": f"{prefix}（选择题）.docx",
        "written": f"{prefix}（主观题）.docx",
        "analysis": f"{prefix} 解析.docx",
    }


def load_config(args):
    path = ASSETS / "defaults.json"
    config = json.loads(path.read_text(encoding="utf-8"))
    for item in getattr(args, "set", []) or []:
        if "=" not in item:
            raise ValueError("配置必须使用key=value")
        key, value = item.split("=", 1)
        if key not in config:
            raise ValueError(f"未知配置项：{key}")
        if key == "target_analysis_pages":
            config[key] = int(value)
        elif key == "analysis_fill_threshold":
            config[key] = float(value)
        else:
            config[key] = value
    if not isinstance(config["target_analysis_pages"], int) or config["target_analysis_pages"] < 1:
        raise ValueError("target_analysis_pages必须为正整数")
    if not isinstance(config["analysis_fill_threshold"], (int, float)) or not 0.5 <= config["analysis_fill_threshold"] <= 1:
        raise ValueError("analysis_fill_threshold必须在0.5至1之间")
    for key, value in config.items():
        if key not in ("target_analysis_pages", "analysis_fill_threshold") and (not isinstance(value, str) or "\n" in value):
            raise ValueError(f"配置{key}必须为单行字符串")
    return path, config


def convert(args, config):
    inspection, question_doc, questions, _, analysed = inspect_pair(Path(args.question), Path(args.analysis))
    if inspection["errors"]:
        print(json.dumps(inspection, ensure_ascii=False, indent=2))
        raise RuntimeError("两份输入未通过一致性检查，转换已停止")
    root = Path(args.out)
    root.mkdir(parents=True, exist_ok=True)
    run = root / ("精品解析转换-" + datetime.now().strftime("%Y%m%d-%H%M%S-%f"))
    outputs = run / "outputs"
    qa = run / "qa"
    internal = run / "internal"
    for directory in (outputs, qa, internal):
        directory.mkdir(parents=True)
    write_json(run / "inspection.json", inspection)
    names = output_names(config)
    records = analysis_records(analysed)
    reference_manifest = json.loads((ASSETS / "reference-manifest.json").read_text(encoding="utf-8"))
    reference_pair = (
        inspection["question_sha256"] == reference_manifest["question_sha256"]
        and inspection["analysis_sha256"] == reference_manifest["analysis_sha256"]
    )
    if reference_pair:
        customize_reference("selection", config, outputs / names["selection"])
        customize_reference("written", config, outputs / names["written"])
        selection_numbers = [q["number"] for q in questions if q["kind"] == "choice"]
        written_numbers = [q["number"] for q in questions if q["kind"] == "written"]
    else:
        selection_numbers = make_question_doc("selection", config, question_doc, questions, outputs / names["selection"])
        written_numbers = make_question_doc("written", config, question_doc, questions, outputs / names["written"])
    written_heading = next((h for h in inspection["headings"].get("written", []) if "综合题" in h), None)
    reference_counts = {int(key): value for key, value in reference_manifest["written_detail_counts"].items()}
    requested_counts = parse_detail_counts(args.detail_count, records) if args.detail_count else reference_counts
    exact_reference_analysis = (
        reference_pair
        and int(config["target_analysis_pages"]) == int(reference_manifest["analysis_pages"])
        and requested_counts == reference_counts
    )
    if exact_reference_analysis:
        analysis_output = outputs / names["analysis"]
        customize_reference("analysis", config, analysis_output)
        best = {
            "label": "exact-reference-template",
            "docx": str(analysis_output),
            "detail_counts": reference_counts,
        }
        trials = [best]
        warnings = []
    elif args.detail_count:
        best, trials, warnings = make_manual_detail_choice(
            config, records, written_heading, internal, args.detail_count
        )
    else:
        best, trials, warnings = choose_detail_fill(config, records, written_heading, internal)
        warnings.append("自动详解范围是LibreOffice辅助排版候选；交付前必须用WPS复核并按需用--detail-count重跑")
    if not exact_reference_analysis:
        shutil.copy2(best["docx"], outputs / names["analysis"])

    rendered = {}
    for key, filename in names.items():
        pdf, pages = render_pdf(outputs / filename, qa / key)
        pngs = render_pngs(pdf, qa / key / "pages")
        rendered[key] = {
            "filename": filename,
            "sha256": digest(outputs / filename),
            "pages": pages,
            "page_images": [str(path.resolve()) for path in pngs],
        }
    manifest = {
        "config": config,
        "selection_questions": selection_numbers,
        "written_questions": written_numbers,
        "analysis_questions": [record["number"] for record in records],
        "format_mode": "exact-reference-template" if reference_pair else "generic-template-reflow",
        "written_detail_counts": best["detail_counts"],
        "detail_fill_trials": trials,
        "warnings": warnings + inspection["review_issues"],
        "files": rendered,
    }
    write_json(run / "manifest.json", manifest)
    print(json.dumps({"run": str(run.resolve()), **manifest}, ensure_ascii=False, indent=2))
    return run


def verify(run: Path):
    manifest_path = run / "manifest.json"
    if not manifest_path.exists():
        raise ValueError("运行目录缺少manifest.json")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    errors = []
    for key, record in manifest["files"].items():
        path = run / "outputs" / record["filename"]
        if not path.exists():
            errors.append(f"缺少输出：{record['filename']}")
            continue
        if digest(path) != record["sha256"]:
            errors.append(f"文件在渲染后发生变化：{record['filename']}")
        if len(record.get("page_images", [])) != record["pages"]:
            errors.append(f"页面图片数量不符：{record['filename']}")
    result = {"run": str(run.resolve()), "errors": errors, "files": manifest["files"]}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if errors:
        raise RuntimeError("输出验证失败")


def main():
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    inspect_parser = sub.add_parser("inspect")
    inspect_parser.add_argument("question")
    inspect_parser.add_argument("analysis")
    inspect_parser.add_argument("--out")
    convert_parser = sub.add_parser("convert")
    convert_parser.add_argument("question")
    convert_parser.add_argument("analysis")
    convert_parser.add_argument("--out", required=True)
    convert_parser.add_argument("--set", action="append", default=[])
    convert_parser.add_argument(
        "--detail-count",
        action="append",
        default=[],
        metavar="题号=段落数",
        help="按WPS复核结果指定各主观题从详解开头连续保留的段落数，可重复使用",
    )
    verify_parser = sub.add_parser("verify")
    verify_parser.add_argument("run")
    config_parser = sub.add_parser("config")
    config_parser.add_argument("--set", action="append", default=[])
    config_parser.add_argument("--save-defaults", action="store_true")
    args = parser.parse_args()
    config_path, config = load_config(args)
    if args.command == "inspect":
        result, *_ = inspect_pair(Path(args.question), Path(args.analysis))
        if args.out:
            write_json(Path(args.out), result)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        if result["errors"]:
            raise RuntimeError("输入一致性检查失败")
    elif args.command == "convert":
        convert(args, config)
    elif args.command == "verify":
        verify(Path(args.run))
    elif args.command == "config":
        if args.save_defaults:
            write_json(config_path, config)
        print(json.dumps(config, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
